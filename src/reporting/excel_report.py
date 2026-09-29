from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter


class ExcelReport:

    def generate(self, violations, output_file):
        try:
            import xlsxwriter
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "Excel export requires XlsxWriter. Install project dependencies "
                "into the same Python environment that runs Streamlit with: "
                "python -m pip install -r requirements.txt"
            ) from exc

        output_path = Path(output_file)

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        # ======================================================
        # Violation DataFrame
        # ======================================================

        if violations:
            normalized = [
                item.to_dict() if hasattr(item, "to_dict") else item
                for item in violations
            ]
            df = pd.DataFrame(normalized)

        else:

            df = pd.DataFrame()

        # ======================================================
        # Report columns
        # ======================================================

        report_columns = [
            "record_id",
            "timestamp",
            "user_id",
            "department",
            "role",
            "action",
            "status",
            "data_type",
            "resource",
            "transaction_amount",
            "approved_amount",
            "approval_limit",
            "approver_id",
            "sod_conflict",
            "consent",
            "justification",
            "data_access_count_24h",
            "ip_address",
            "rule_id",
            "rule_name",
            "description",
            "severity",
            "field",
            "operator",
            "actual_value",
            "expected_value",
            "source_row",
            "detected_at"
        ]

        # Ensure all columns exist
        for column in report_columns:

            if column not in df.columns:
                df[column] = ""

        df = df[report_columns]

        # ======================================================
        # Friendly column names
        # ======================================================

        column_mapping = {

            "record_id": "Record ID",
            "timestamp": "Event Timestamp",
            "user_id": "User ID",
            "department": "Department",
            "role": "Role",
            "action": "Action",
            "status": "Status",
            "data_type": "Data Type",
            "resource": "Resource",

            "transaction_amount": "Transaction Amount",
            "approved_amount": "Approved Amount",
            "approval_limit": "Approval Limit",
            "approver_id": "Approver ID",

            "sod_conflict": "SoD Conflict",
            "consent": "Consent",
            "justification": "Justification",
            "data_access_count_24h": "Data Access / 24h",

            "ip_address": "IP Address",

            "rule_id": "Rule ID",
            "rule_name": "Rule Name",
            "description": "Rule Description",
            "severity": "Severity",

            "field": "Evaluated Field",
            "operator": "Operator",
            "actual_value": "Actual Value",
            "expected_value": "Expected Value",

            "source_row": "Source CSV Row",
            "detected_at": "Detected At"
        }

        df = df.rename(
            columns=column_mapping
        )

        # ======================================================
        # Summary
        # ======================================================

        total_violations = len(df)

        if total_violations:

            severity_summary = (
                df["Severity"]
                .astype(str)
                .str.upper()
                .value_counts()
                .reset_index()
            )

            severity_summary.columns = [
                "Severity",
                "Count"
            ]

            rule_summary = (
                df["Rule Name"]
                .value_counts()
                .reset_index()
            )

            rule_summary.columns = [
                "Rule",
                "Count"
            ]

            department_summary = (
                df["Department"]
                .value_counts()
                .reset_index()
            )

            department_summary.columns = [
                "Department",
                "Count"
            ]

            user_summary = (
                df["User ID"]
                .value_counts()
                .reset_index()
            )

            user_summary.columns = [
                "User ID",
                "Violation Count"
            ]

        else:

            severity_summary = pd.DataFrame(
                columns=["Severity", "Count"]
            )

            rule_summary = pd.DataFrame(
                columns=["Rule", "Count"]
            )

            department_summary = pd.DataFrame(
                columns=["Department", "Count"]
            )

            user_summary = pd.DataFrame(
                columns=["User ID", "Violation Count"]
            )

        # ======================================================
        # Write Excel
        # ======================================================

        try:

            dashboard = pd.DataFrame({
                "Metric": ["Total Violations", "Critical Violations", "High Violations", "Medium Violations", "Low Violations"],
                "Value": [
                    total_violations,
                    self._count_severity(df, "CRITICAL"),
                    self._count_severity(df, "HIGH"),
                    self._count_severity(df, "MEDIUM"),
                    self._count_severity(df, "LOW"),
                ],
            })
            # XlsxWriter's constant-memory mode streams rows as it writes.
            # This avoids the large cell-object and XML overhead of openpyxl.
            sheets = [
                ("Dashboard", dashboard),
                ("Violations", df),
                ("By Severity", severity_summary),
                ("By Rule", rule_summary),
                ("By Department", department_summary),
                ("By User", user_summary),
            ]
            workbook = xlsxwriter.Workbook(
                str(output_path),
                {"constant_memory": True, "strings_to_urls": False},
            )
            header_format = workbook.add_format({"bold": True, "align": "center", "valign": "vcenter"})
            for sheet_name, data in sheets:
                worksheet = workbook.add_worksheet(sheet_name)
                worksheet.freeze_panes(1, 0)
                if len(data.columns):
                    worksheet.autofilter(0, 0, len(data), len(data.columns) - 1)
                    worksheet.write_row(0, 0, list(data.columns), header_format)
                    for column_index, column_name in enumerate(data.columns):
                        worksheet.set_column(
                            column_index,
                            column_index,
                            min(max(len(str(column_name)) + 2, 12), 45),
                        )

                for row_index, values in enumerate(data.itertuples(index=False, name=None), start=1):
                    row = []
                    for value in values:
                        try:
                            if pd.isna(value):
                                value = None
                        except (TypeError, ValueError):
                            pass
                        if hasattr(value, "to_pydatetime"):
                            value = value.to_pydatetime()
                        elif hasattr(value, "item"):
                            value = value.item()
                        row.append(value)
                    worksheet.write_row(row_index, 0, row)

            if len(df):
                violations_sheet = workbook.get_worksheet_by_name("Violations")
                severity_index = list(df.columns).index("Severity")
                severity_letter = get_column_letter(severity_index + 1)
                critical_format = workbook.add_format({"bg_color": "#FFC7CE"})
                high_format = workbook.add_format({"bg_color": "#FFEB9C"})
                violations_sheet.conditional_format(
                    1, severity_index, len(df), severity_index,
                    {"type": "formula", "criteria": f'=${severity_letter}2="CRITICAL"', "format": critical_format},
                )
                violations_sheet.conditional_format(
                    1, severity_index, len(df), severity_index,
                    {"type": "formula", "criteria": f'=${severity_letter}2="HIGH"', "format": high_format},
                )
            workbook.close()

            print(
                "Excel compliance report generated successfully:"
            )

            print(
                output_path.resolve()
            )

        except PermissionError:

            raise PermissionError(
                f"Unable to write Excel report: "
                f"{output_path.resolve()}. "
                f"Close the Excel file if it is open."
            )

    # ==========================================================
    # Formatting
    # ==========================================================

    def _format_workbook(self, output_path):

        from openpyxl import load_workbook

        workbook = load_workbook(
            output_path
        )

        for worksheet in workbook.worksheets:

            worksheet.freeze_panes = "A2"

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

            # Header
            for cell in worksheet[1]:

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            # Column width
            # Widths are capped at 45 characters, so sampling the first
            # 1,000 rows gives a useful width without scanning every cell in
            # potentially very large violation reports.
            sample_rows = min(worksheet.max_row, 1000)
            for column_index in range(1, worksheet.max_column + 1):
                max_length = 0
                column_letter = worksheet.cell(1, column_index).column_letter
                for row_index in range(1, sample_rows + 1):
                    try:
                        value_length = len(str(worksheet.cell(row_index, column_index).value))
                        max_length = max(
                            max_length,
                            value_length
                        )
                    except Exception:
                        pass

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max(max_length + 2, 12),
                    45
                )

        # ======================================================
        # Conditional formatting for severity
        # ======================================================

        if "Violations" in workbook.sheetnames:

            worksheet = workbook[
                "Violations"
            ]

            severity_column = None

            for cell in worksheet[1]:

                if cell.value == "Severity":
                    severity_column = cell.column_letter
                    break

            if severity_column:

                last_row = worksheet.max_row

                critical_rule = FormulaRule(
                    formula=[
                        f'{severity_column}2="CRITICAL"'
                    ],
                    fill=PatternFill(
                        fill_type="solid",
                        fgColor="FFC7CE"
                    )
                )

                high_rule = FormulaRule(
                    formula=[
                        f'{severity_column}2="HIGH"'
                    ],
                    fill=PatternFill(
                        fill_type="solid",
                        fgColor="FFEB9C"
                    )
                )

                worksheet.conditional_formatting.add(
                    f"{severity_column}2:"
                    f"{severity_column}{last_row}",
                    critical_rule
                )

                worksheet.conditional_formatting.add(
                    f"{severity_column}2:"
                    f"{severity_column}{last_row}",
                    high_rule
                )

        workbook.save(
            output_path
        )

    # ==========================================================
    # Utility
    # ==========================================================

    @staticmethod
    def _count_severity(df, severity):

        if df.empty:
            return 0

        return int(
            (
                df["Severity"]
                .astype(str)
                .str.upper()
                == severity
            ).sum()
        )
